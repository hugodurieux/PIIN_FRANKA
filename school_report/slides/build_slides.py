#!/usr/bin/env python3
"""Build the 15-slide defence deck as a real .pptx file.

Usage
-----
    ./.venv/bin/pip install python-pptx
    ./.venv/bin/python school_report/slides/build_slides.py

Writes ``school_report/slides/pinn_franka_defence.pptx`` (16:9, 15 slides,
speaker notes with a time budget on every slide).

The equations and the architecture diagram are pre-rendered PNGs in
``assets/``, produced from the same TikZ/LaTeX source as the report, so the
slides carry real typeset maths rather than ASCII approximations. Regenerate
them only if the report's figure changes.

NUMBERS: every figure on slides 11-14 is traceable to a run in results/.
See the provenance note at the bottom of this file.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "pinn_franka_defence.pptx"

# --- palette -----------------------------------------------------------
NAVY = RGBColor(0x0F, 0x3D, 0x62)      # titles, rules, emphasis
INK = RGBColor(0x1A, 0x1A, 0x1A)       # body text
MUTED = RGBColor(0x5A, 0x6B, 0x7A)     # secondary text, footers
ACCENT = RGBColor(0xB0, 0x3A, 0x2E)    # the one "but" per slide
PALE = RGBColor(0xF2, 0xF5, 0xF8)      # table banding
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xC7, 0xD2, 0xDC)

FONT = "Calibri"

# --- geometry (13.333 x 7.5 in, 16:9) ----------------------------------
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.75)
BODY_W = W - 2 * MARGIN
TITLE_TOP = Inches(0.45)
BODY_TOP = Inches(1.55)
BODY_BOTTOM = Inches(6.85)


def _para(p, text, size, *, bold=False, color=INK, italic=False,
          space_after=8, align=PP_ALIGN.LEFT, level=0):
    p.text = text
    p.level = level
    p.alignment = align
    p.space_after = Pt(space_after)
    f = p.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return p


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def bullets(slide, items, *, top=BODY_TOP, left=MARGIN, width=BODY_W,
            size=21, height=None):
    """items: list of (text, level[, style]); style in {'', 'b', 'accent',
    'muted'}."""
    tf = textbox(slide, left, top, width, height or (BODY_BOTTOM - top))
    for i, item in enumerate(items):
        text, level, style = (item + ("",))[:3] if len(item) < 3 else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        color = {"accent": ACCENT, "muted": MUTED}.get(style, INK)
        prefix = "" if level == 0 else "– "
        _para(p, prefix + text, size - 2 * level, level=level,
              bold=(style == "b"), color=color,
              space_after=10 if level == 0 else 6)
    return tf


def slide_shell(prs, title, kicker=None, number=None):
    """Blank slide with the standard title block, rule and footer."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if kicker:
        tf = textbox(slide, MARGIN, TITLE_TOP - Inches(0.02), BODY_W,
                     Inches(0.3))
        _para(tf.paragraphs[0], kicker.upper(), 12, bold=True, color=MUTED,
              space_after=0)
        t_top = TITLE_TOP + Inches(0.3)
    else:
        t_top = TITLE_TOP

    tf = textbox(slide, MARGIN, t_top, BODY_W, Inches(0.8))
    _para(tf.paragraphs[0], title, 30, bold=True, color=NAVY, space_after=0)

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.34), BODY_W, Emu(12700))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.fill.background()
    line.shadow.inherit = False

    if number is not None:
        tf = textbox(slide, W - MARGIN - Inches(1.2), Inches(6.95),
                     Inches(1.2), Inches(0.3))
        _para(tf.paragraphs[0], f"{number} / 15", 11, color=MUTED,
              align=PP_ALIGN.RIGHT, space_after=0)
    return slide


def picture(slide, name, *, top, width, left=None):
    """Insert an asset scaled to `width`, aspect ratio preserved."""
    path = ASSETS / name
    if not path.exists():
        raise SystemExit(f"missing asset: {path}")
    pic = slide.shapes.add_picture(str(path), Inches(0), top, width=width)
    pic.left = left if left is not None else int((W - pic.width) / 2)
    return pic


def caption(slide, text, top, *, size=14):
    tf = textbox(slide, MARGIN, top, BODY_W, Inches(0.4))
    _para(tf.paragraphs[0], text, size, color=MUTED, italic=True,
          align=PP_ALIGN.CENTER, space_after=0)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def table(slide, rows, *, top, left=MARGIN, width=BODY_W, col_w=None,
          size=15, header_size=14, row_h=Inches(0.42),
          highlight_row=None, rule_before=()):
    """rows[0] is the header. col_w: list of Inches summing to `width`.
    rule_before: row indices that get a top border (a booktabs-style midrule
    is not available, so a paler fill marks the separation instead)."""
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width,
                                   row_h * n_rows)
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False

    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = cw
    for r in range(n_rows):
        tbl.rows[r].height = row_h

    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
                color, bold, sz = WHITE, True, header_size
            elif r == highlight_row:
                cell.fill.fore_color.rgb = RGBColor(0xE3, 0xEC, 0xF4)
                color, bold, sz = NAVY, True, size
            elif r in rule_before:
                cell.fill.fore_color.rgb = RGBColor(0xE9, 0xED, 0xF1)
                color, bold, sz = MUTED, True, size
            else:
                cell.fill.fore_color.rgb = PALE if r % 2 else WHITE
                color, bold, sz = INK, False, size
            p = cell.text_frame.paragraphs[0]
            _para(p, str(val), sz, bold=bold, color=color, space_after=0,
                  align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    return tbl


# =======================================================================
#  Slides
# =======================================================================

def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # --- 1. Title ------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                              Inches(0.22), H)
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False

    tf = textbox(s, Inches(1.1), Inches(2.0), Inches(11.2), Inches(2.4))
    _para(tf.paragraphs[0],
          "Grey-box physics-informed neural networks for "
          "payload-conditioned dynamics learning and control of a 7 DoF "
          "manipulator",
          36, bold=True, color=NAVY, space_after=14)
    _para(tf.add_paragraph(),
          "An automated pipeline from URDF file to controlled robot, and "
          "1 kHz real-time control of a simulated Franka Emika Panda "
          "(Isaac Sim / MuJoCo)",
          19, color=MUTED, space_after=0)

    tf = textbox(s, Inches(1.1), Inches(5.4), Inches(11.2), Inches(1.2))
    _para(tf.paragraphs[0], "Hugo Durieux", 20, bold=True, color=INK,
          space_after=4)
    _para(tf.add_paragraph(),
          "Master's research internship  ·  University of Melbourne", 15,
          color=MUTED, space_after=0)
    notes(s, """
[0:00 – 0:20]  Title slide.
Say who you are and the one-sentence framing: the input is a URDF file, the
output is a robot you can control in real time, and the contribution sits in
the dynamics model in between.
Do not read the title aloud.
""")

    # --- 2. Motivation -------------------------------------------------
    s = slide_shell(prs, "A model that is exact — and still wrong",
                    kicker="Motivation", number=2)
    bullets(s, [
        ("Torque-level control needs a faithful inverse dynamics model.", 0),
        ("The URDF gives one for free: geometry and inertia, exactly, and "
         "RNEA evaluates it with no manual derivation.", 0),
        ("But the rigid-body model knows nothing about:", 0),
        ("dry and viscous friction in the harmonic drives", 1),
        ("transmission compliance and mechanical backlash", 1),
        ("temperature-dependent losses, actuator wear, the carried payload", 1),
        ("The gap is several newton-metres on a Panda, and it is not "
         "analytically modellable. That gap is the whole subject.", 0,
         "accent"),
    ], size=22)
    notes(s, """
[0:20 – 1:20]  The problem, in one idea.
The analytical model is not approximately right and slightly noisy — it is
structurally incapable of representing the effects that dominate at the joint
level. Emphasise: we are not replacing physics, we are completing it. That
framing carries the whole talk.
""")

    # --- 3. Objective and the three difficulties -----------------------
    s = slide_shell(prs, "Objective, and why it is not just engineering",
                    kicker="Positioning", number=3)
    bullets(s, [
        ("Goal: an automatic pipeline whose only input is a description file "
         "and whose output is a robot effectively controlled by a learned "
         "model.", 0, "b"),
        ("Positioned against Liu, Borja & Della Santina (2024): same robot, "
         "same method class, same purpose.", 0, "muted"),
        ("Three difficulties define the contribution:", 0),
        ("Generality costs accuracy — no hand-derived equations, everything "
         "extracted from the URDF.", 1),
        ("Physical guarantees survive learning badly — 7 joints whose torque "
         "limits differ by a factor of 7 (87 vs 12 N·m).", 1),
        ("Offline accuracy is not control — sub-millisecond inference, and a "
         "control law whose error can be bounded.", 1),
    ], size=21)
    notes(s, """
[1:20 – 2:20]  What is claimed, and what is not.
Four contributions in the report; these three difficulties are why it is not a
plumbing exercise. State plainly what is NOT claimed: no new learning theory,
no hardware — Isaac Sim for training, MuJoCo for deployment.
""")

    # --- 4. Preliminaries ----------------------------------------------
    s = slide_shell(prs, "What the URDF gives for free",
                    kicker="Preliminaries", number=4)
    picture(s, "eq_rbd.png", top=Inches(1.95), width=Inches(8.4))
    bullets(s, [
        ("M(q) inertia · C(q, q̇) Coriolis · G(q) gravity, n = 7 for the "
         "Panda.", 0),
        ("Evaluated analytically by the recursive Newton–Euler algorithm "
         "(RNEA), through Pinocchio, straight from the URDF.", 0),
        ("Never learned, never modified — this term is the white box, and "
         "nothing in the training loop can touch it.", 0, "b"),
    ], top=Inches(3.4), size=20)
    notes(s, """
[2:20 – 3:10]  Standard, move fast.
The audience knows this equation. The only point worth landing: RNEA is
evaluated, not fitted. That is what makes the guarantees downstream mean
something, and it is what makes the pipeline automatic — any URDF, no
re-derivation.
""")

    # --- 5. The grey-box decomposition ---------------------------------
    s = slide_shell(prs, "Computed, plus learned", kicker="Method",
                    number=5)
    picture(s, "eq_model.png", top=Inches(1.9), width=Inches(11.0))
    bullets(s, [
        ("The residual absorbs dissipation and modelling error only.", 0),
        ("Conditioned on the carried payload δ — following Hu et al., "
         "backlash error depends on position and velocity but not on mass, "
         "so only the residual is conditioned.", 0),
        ("τ_res does not depend on the acceleration q̈: available offline, "
         "never inside a 1 kHz loop that measures only q and q̇.", 0, "b"),
    ], top=Inches(3.6), size=20)
    notes(s, """
[3:10 – 4:10]  The core idea. Take your time.
Point at the two braces. Then the real-time argument: excluding q̈ from the
residual is a deployment requirement, not a modelling preference. It comes
back on slide 12 — it is also why the black-box result is suspicious.
""")

    # --- 6. Constraints ------------------------------------------------
    s = slide_shell(prs, "Three degrees of guarantee", kicker="Method",
                    number=6)
    picture(s, "eq_problem.png", top=Inches(1.8), width=Inches(11.2))
    bullets(s, [
        ("Degree 0 — free: a smooth, payload-conditioned residual "
         "(Mish, C¹ continuous; never ReLU).", 0),
        ("Degree 1 — imposed: augmented Lagrangian with dual ascent, so the "
         "multipliers are optimised instead of ρ being driven to infinity.",
         0),
        ("Degree 2 — structural: guaranteed by construction, next slide but "
         "one.", 0, "b"),
    ], top=Inches(3.7), size=19)
    notes(s, """
[4:10 – 5:10]  The constraints are the contribution, not the loss.
Two constraints: never command a torque the motor cannot deliver, and never
let the learned term inject energy — the formal translation of "this term
represents friction". The three-degree ladder is the structure of the whole
method section; announce it here and refer back to it twice.
""")

    # --- 7. Architecture -----------------------------------------------
    s = slide_shell(prs, "Architecture and gradient flow",
                    kicker="Method", number=7)
    picture(s, "pipeline.png", top=Inches(1.7), width=Inches(9.3))
    caption(s, "Grey background: computed from the URDF, no learning. "
               "Thick outline: trained. Dashed: losses and backpropagation.",
            Inches(6.45))
    notes(s, """
[5:10 – 6:10]  One figure, one path through it.
Trace it left to right, once: URDF → RNEA → τ_rbd, and separately measured
state → encoding → the two residual networks. They meet at the sum. Then the
losses along the bottom and the single gradient path — the gradient reaches
only the two trained networks, never RNEA.
""")

    # --- 8. FrictionNet ------------------------------------------------
    s = slide_shell(prs, "Degree 2: dissipativity by construction",
                    kicker="Method", number=8)
    picture(s, "eq_fric.png", top=Inches(1.9), width=Inches(8.6))
    bullets(s, [
        ("D positive-definite and diagonal ⇒ τ_fric·q̇ ≤ 0 for every input.",
         0, "b"),
        ("The guarantee depends neither on the data nor on the convergence "
         "of the optimisation — it is a property of the function class.", 0),
        ("Honest caveat: the guarantee bears on the sign, not the magnitude. "
         "A strongly negative weight tied to a large δ could drive the "
         "friction towards zero.", 0, "muted"),
    ], top=Inches(3.3), size=20)
    notes(s, """
[6:10 – 7:00]  Why this matters more than accuracy.
This is the property the black box cannot offer, and on slide 13 you will show
that FrictionNet buys only 1.4 % of accuracy. Say it here first, so that the
ablation reads as a confirmation of your framing rather than as a
disappointment: the module is there for the guarantee, not for the RMSE.
""")

    # --- 9. Closing the loop -------------------------------------------
    s = slide_shell(prs, "From model to torque, at 1 kHz",
                    kicker="Control", number=9)
    picture(s, "eq_control.png", top=Inches(1.75), width=Inches(10.2))
    tf = textbox(s, MARGIN, Inches(2.75), BODY_W, Inches(0.5))
    _para(tf.paragraphs[0],
          "Gains from the 99.9 % error quantile of the test set (ε_j), "
          "following Liu et al.:", 19, color=INK, space_after=0)
    picture(s, "eq_gains.png", top=Inches(3.4), width=Inches(6.0))
    bullets(s, [
        ("A counter-intuitive consequence: K_p derives from ε_j, so a "
         "better-modelled joint gets a lower gain — and e_ss = τ / K_p means "
         "the best-learned joints track worst.", 0, "accent"),
        ("Measured live on joint 5. This is a limitation of the template, "
         "not an implementation defect.", 0, "muted"),
    ], top=Inches(4.6), size=19)
    notes(s, """
[7:00 – 8:00]  Closing the loop.
ε_j is measured, not tuned: a better model gives a smaller bound, lower gains,
a softer arm. Then deliver the counter-intuitive consequence deliberately — it
shows you understood the template you borrowed rather than just applying it,
and it is confirmed by a measurement later in the talk.
""")

    # --- 10. Experimental design ---------------------------------------
    s = slide_shell(prs, "Experimental design", kicker="Methods", number=10)
    bullets(s, [
        ("Data: 30 continuous 5 s trajectories in Isaac Sim, Fourier "
         "excitation around Sobol centres, three payloads, 1 kHz → 148,304 "
         "samples.", 0, "b"),
        ("Clipped samples filtered out (|τ| ≤ 0.97 τ̄): a saturated sample "
         "reports the actuator bound, not the dynamics.", 1),
        ("80 / 10 / 10 split, shared by training, evaluation and ε_j.", 1),
        ("Three models, same test set, same capacity, same 200 epochs:", 0,
         "b"),
        ("RNEA alone — the URDF, no learning.", 1),
        ("Direct MLP — same input, same capacity, no τ_rbd term.", 1),
        ("Grey box — τ_rbd + constrained residual.", 1),
        ("Until 31 July 2026 the project had no test set at all: the reported "
         "RMSE was the validation one, on the set the checkpoint was selected "
         "on. These are the first figures measured outside that.", 0,
         "accent"),
    ], size=19)
    notes(s, """
[8:00 – 9:00]  Protocol. Be crisp — this is a setup slide, but the last line
is not a detail.
Owning the missing-test-set finding yourself is what makes everything after it
credible. It also pre-empts the harshest question available to the examiner.
""")

    # --- 11. Results: baselines ----------------------------------------
    s = slide_shell(prs, "Prediction error against the baselines",
                    kicker="Results", number=11)
    table(s, [
        ["Model", "Mean RMSE\n[N·m]", "Worst joint\n[N·m]",
         "Over torque\nlimit", "Guarantees"],
        ["RNEA alone (analytical)", "1.456", "2.841", "147  (0.99 %)",
         "none (not trained)"],
        ["Direct MLP (black box)", "0.447", "0.971", "0  (0.00 %)",
         "torque limits only"],
        ["Grey box (proposed)", "0.624", "1.022", "0  (0.00 %)",
         "limits + dissipativity"],
    ], top=Inches(1.85), row_h=Inches(0.62),
        col_w=[Inches(3.6), Inches(1.9), Inches(1.9), Inches(2.0),
               Inches(2.44)],
        highlight_row=2)
    bullets(s, [
        ("The grey box cuts the analytical model's error by 2.3× and holds "
         "zero torque-limit violations where RNEA violates on 0.99 % of "
         "samples.", 0),
        ("But a black box of identical capacity and budget is 28 % more "
         "accurate. This contradicts the starting hypothesis and is reported "
         "as a result, not hidden.", 0, "accent"),
    ], top=Inches(4.6), size=19)
    notes(s, """
[9:00 – 10:15]  The headline, and it is not the one you set out to get.
Read the accuracy column, then the guarantees column — the one the black box
cannot fill in. Do not apologise for the black-box row; announce it, then say
"and the next slide explains why that number is an upper bound". That
sequencing turns the weakest moment of the talk into the strongest.
""")

    # --- 12. The partition caveat --------------------------------------
    s = slide_shell(prs, "A caveat on the partition — and what it distorts",
                    kicker="Results", number=12)
    bullets(s, [
        ("The data are 30 continuous trajectories at 1 kHz, but the split is "
         "drawn per sample.", 0, "b"),
        ("Two consecutive steps, 1 ms apart on the same trajectory, land on "
         "opposite sides: every test point has near-identical neighbours in "
         "training.", 1),
        ("The selection defect is fixed. Sample independence is not.", 0),
        ("The leakage does not help the three models equally — it helps most "
         "the model that has to memorise most.", 0, "b"),
        ("The clue: the black box never receives q̈, so it cannot represent "
         "the inertial torque M(q)q̈, the dominant term — and it still wins "
         "by 28 %. With 30 trajectories, (q, q̇) almost always identifies the "
         "trajectory and phase, making q̈ recoverable by memorisation.", 1,
         "accent"),
        ("So the 28 % is an upper bound on the black box's advantage, not an "
         "established result. The correct measurement is a group-wise split "
         "by trajectory — implemented, not yet reportable.", 0),
    ], size=19)
    notes(s, """
[10:15 – 11:30]  The slide that shows you can criticise your own result.
This is a methodological argument, so deliver it as one: state the mechanism,
then the falsifiable clue (no q̈, yet it wins), then the correction. Do not
overclaim — say honestly that the group-wise split exists in the repo but
rests on too few test trajectories to report. If asked what it showed
preliminarily: the black box collapses, which is what the argument predicts.
""")

    # --- 13. Per-joint and ablations -----------------------------------
    s = slide_shell(prs, "The residual degrades what RNEA already gets right",
                    kicker="Results", number=13)
    table(s, [
        ["Test RMSE [N·m]", "J1", "J2", "J3", "J4", "J5", "J6", "J7"],
        ["RNEA alone", "1.588", "2.841", "1.824", "1.953", "0.883", "1.062",
         "0.042"],
        ["Direct MLP", "0.971", "0.845", "0.405", "0.742", "0.066", "0.068",
         "0.031"],
        ["Grey box (proposed)", "1.022", "0.935", "0.553", "0.843", "0.371",
         "0.355", "0.288"],
    ], top=Inches(1.8), row_h=Inches(0.5),
        col_w=[Inches(3.3), Inches(1.12), Inches(1.12), Inches(1.12),
               Inches(1.12), Inches(1.12), Inches(1.12), Inches(1.12)])
    bullets(s, [
        ("J7 is a wrist roll: RNEA reaches 0.042 N·m, 0.35 % of its limit — "
         "there is nothing to learn. The residual makes it 6.9× worse.", 0,
         "accent"),
        ("Mechanism, from the worst samples: τ_real = −0.01, τ_rbd = −0.08 "
         "(already correct), and the network emits τ_res = −2.04. Learned "
         "D₇₇ converges to 0.67, the largest of all seven joints, while "
         "D₁₁ collapses to 0.019.", 0),
        ("Ablations, same protocol: sin/cos encoding buys 1.1 %, FrictionNet "
         "buys 1.4 %. Neither is justified by accuracy.", 0, "muted"),
    ], top=Inches(4.1), size=18)
    notes(s, """
[11:30 – 12:45]  Report the negative result yourself, with its mechanism.
The unweighted MSE aggregates over joints, so a joint contributing almost
nothing to the total gets almost no gradient pressure to stay at zero. The
D₇₇ figure is the evidence that this is real and not a fluke: the dissipative
module is most active exactly where there is no friction to model.
The fix is named in the report — a diag(1/τ̄²) weighting.
""")

    # --- 14. Closed-loop ablation --------------------------------------
    s = slide_shell(prs, "Closed loop: the cost of cross-simulator transfer",
                    kicker="Results", number=14)
    table(s, [
        ["Run", "Residual", "Joint 5 bias [rad]", "Flange error [m]"],
        ["36", "on", "−0.0264", "0.0141"],
        ["37", "on", "−0.0434", "0.0139"],
        ["38", "on", "−0.0312", "0.0149"],
        ["45", "off", "−0.0011", "0.0076"],
        ["46", "off", "−0.0019", "0.0073"],
        ["mean", "on → off", "−0.0337 → −0.0015   (22×)",
         "0.0143 → 0.0075   (2×)"],
    ], top=Inches(1.8), row_h=Inches(0.44),
        col_w=[Inches(1.3), Inches(2.0), Inches(4.3), Inches(4.23)],
        highlight_row=6, rule_before=(4, 5))
    bullets(s, [
        ("Trained on Isaac Sim, deployed under MuJoCo: the residual learned "
         "the gap between RNEA and *Isaac's* dynamics, and adds it to an "
         "engine that does not have it.", 0),
        ("Defensible: the Isaac-trained residual degrades tracking under "
         "cross-simulator transfer. NOT: that the grey-box approach fails.",
         0, "b"),
        ("Confounder ruled out — the gain margin is held constant, and a "
         "lower gain would push e_ss the other way.", 0, "muted"),
    ], top=Inches(5.0), size=18)
    notes(s, """
[12:45 – 13:50]  Two independent measurements, one cause.
Tie it back to slide 13 explicitly: offline, the residual degrades joint 7 by
6.9×; in closed loop, disabling it improves joint 5 by 22×. Same phenomenon,
measured two different ways.
Then frame it as an asset: this is the quantified baseline that the
frozen-backbone fine-tuning protocol has to beat.
""")

    # --- 15. Conclusion ------------------------------------------------
    s = slide_shell(prs, "What the grey structure is actually for",
                    kicker="Wrap-up", number=15)
    bullets(s, [
        ("A URDF file in, a payload-conditioned dynamics model out, driving a "
         "real 1 kHz control loop — with torque limits and dissipativity "
         "guaranteed rather than hoped for.", 0, "b"),
        ("The case for the grey box is NOT accuracy. On these data a black "
         "box beats it. The case is what the black box cannot offer:", 0,
         "accent"),
        ("guarantees by construction, not by penalty", 1),
        ("an automatic pipeline that never relearns the inertia and gravity "
         "the URDF already provides", 1),
        ("Next, by decreasing scientific value:", 0, "b"),
        ("run the frozen-backbone fine-tuning and replay the ablation above", 1),
        ("a group-wise split by trajectory, to settle the 28 %", 1),
        ("exercise payload conditioning by varying the grasped mass", 1),
        ("Thank you — questions welcome.", 0, "muted"),
    ], size=18)
    notes(s, """
[13:50 – 15:00]  Land it, then stop talking.
The reframing is the point of the whole talk: you set out to show the grey box
is more accurate, you measured that it is not, and you now argue the case on
the grounds that survive. Do not introduce anything new.
Likely questions: why not the real robot (no hardware access during the
internship); why grey box at all (this slide); how the split was made
(slide 12); what happens to J7 (slide 13); what the group-wise split showed
(directionally, the black box collapses — say it is not yet reportable).
""")

    prs.save(OUT)
    print(f"wrote {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()

# -----------------------------------------------------------------------
# PROVENANCE OF THE NUMBERS
#
# Slides 11 and 13   results/20260731_164630/p5_comparison.json
#                    (sample-wise split, seed 12345, 14,830 test samples)
#                    and school_report/rapport/main.tex tables 1-2.
# Slide 13 ablations results/20260731_164630/ phases 4 and 7
#                    (0.631 raw encoding, 0.633 no FrictionNet).
# Slide 14           SESSION.md 2026-07-29 phase F, runs 36/37/38/45/46;
#                    same table as main.tex table 5.
# Slide 12           the caveat of main.tex section 6.2. The group-wise
#                    split referred to lives in results/segsplit_20260803_140610/
#                    (RNEA 1.185, grey box 1.244, black box 4.481 N·m, means
#                    over seeds 12345/23456/34567). It is deliberately NOT on
#                    a slide: three test trajectories per seed is too few to
#                    report, exactly as the report says.
# -----------------------------------------------------------------------
